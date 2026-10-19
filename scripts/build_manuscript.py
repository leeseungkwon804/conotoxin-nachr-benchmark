#!/usr/bin/env python3
"""
Build the Marine Drugs submission .docx from DRAFT_v1.md, inside the journal's own template.

Why not pandoc: `--reference-doc` imports the template's style *definitions* but still writes
pandoc's generic style names (Heading 1, Body Text, ...). The Marine Drugs template defines only
MDPI-specific names (MDPI_2.1_heading1, MDPI_3.1_text, MDPI_5.1_figure_caption, ...) and has no
generic "Heading 1" at all, so every paragraph lands as Normal and the journal formatting is lost.

Why not "clear the body and rebuild it" either: the template's front matter carries pieces that
have no markdown counterpart and that the journal expects verbatim — the article-type line, the
Academic Editor / submission-history table, the affiliation and correspondence placeholders, the
rule under the keywords, the abbreviations scaffold and the publisher's disclaimer. Rebuilding
drops them.

So this edits the template in place: template paragraphs whose content we own get their text
replaced, the template's tutorial sections are deleted, and manuscript content is inserted at the
matching positions. Everything else is left exactly as MDPI shipped it.

Positions are found by matching the template's own placeholder text, not by fixed indices, so an
updated .dot fails loudly rather than silently writing content into the wrong place.

The .dot declares its main part as a Word *template*, which python-docx refuses to open; the copy
made here rewrites that one declaration and touches nothing else.

Usage: python3 scripts/build_manuscript.py [out.docx]
"""
import os
import re
import sys
import zipfile

from PIL import Image
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.text.paragraph import Paragraph

ROOT = "/Volumes/dkyoo_SSD1/review_partime/conotoxin_project"
MD = os.path.join(ROOT, "manuscript", "DRAFT_v1.md")
TEMPLATE = os.path.join(ROOT, "manuscript", "marinedrugs-template.dot")
FIGDIR = os.path.join(ROOT, "manuscript")
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "manuscript", "DRAFT_v3_marinedrugs.docx")

S_TITLE = "MDPI_1.2_title"
S_ABSTRACT = "MDPI_1.7_abstract"
S_KEYWORDS = "MDPI_1.8_keywords"
S_H1 = "MDPI_2.1_heading1"
S_H2 = "MDPI_2.2_heading2"
S_H3 = "MDPI_2.3_heading3"
S_TEXT = "MDPI_3.1_text"
S_TEXT_NOIND = "MDPI_3.2_text_no_indent"
S_FIGCAP = "MDPI_5.1_figure_caption"
S_FIG = "MDPI_5.2_figure"
S_REFS = "MDPI_8.1_references"
S_BACK = "MDPI_6.2_back_matter"

# expanded in the Abbreviations table; order is the order they first appear in the text
ABBREVIATIONS = [
    ("nAChR", "Nicotinic acetylcholine receptor"),
    ("AChBP", "Acetylcholine-binding protein"),
    ("ECD", "Extracellular domain"),
    ("TMD", "Transmembrane domain"),
    ("RMSD", "Root-mean-square deviation"),
    ("PDB", "Protein Data Bank"),
    ("GSO", "Glowworm swarm optimisation"),
    ("MSA", "Multiple sequence alignment"),
    ("fnat", "Fraction of native contacts recovered"),
    ("DockQ", "Continuous docking quality score"),
    ("CAPRI", "Critical Assessment of PRediction of Interactions"),
]

# The A4 text block is 7.27 x 10.08 in. Figures are placed at FIG_W unless that would make them
# taller than FIG_H_MAX, in which case height governs — a figure plus its caption has to fit on one
# page, and two of the seven are tall enough that a fixed width would push the caption over.
FIG_W = 6.9
FIG_H_MAX = 8.3

BACK_MATTER_LABELS = ("Supplementary Materials", "Author Contributions", "Funding",
                      "Institutional Review Board Statement", "Informed Consent Statement",
                      "Data Availability Statement", "Acknowledgments", "Acknowledgements",
                      "Conflicts of Interest")


# ---------------------------------------------------------------- markdown parsing

INLINE = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)", re.S)


def add_runs(par, text):
    """Write text into a paragraph, honouring **bold**, *italic* and `code` spans."""
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            par.add_run(piece[2:-2]).bold = True
        elif piece.startswith("*") and piece.endswith("*"):
            par.add_run(piece[1:-1]).italic = True
        elif piece.startswith("`") and piece.endswith("`"):
            par.add_run(piece[1:-1])
        else:
            par.add_run(piece)


def parse_markdown():
    """-> dict with title, abstract, keywords, body (list of blocks), back (list), refs (list).

    A block is ('image', path) or (style_name, text).
    """
    text = open(MD, encoding="utf-8").read()
    text = re.sub(r"^<!--.*?-->\s*", "", text, flags=re.S)   # working-notes header

    out = {"title": None, "abstract": None, "keywords": None, "body": [], "back": [], "refs": []}
    section = None          # None -> front matter, then 'abstract' / 'body' / 'refs'

    for raw in text.split("\n"):
        line = raw.strip()
        if not line or line == "---":
            continue

        if line.startswith("# "):
            out["title"] = line[2:].strip()
            continue

        if line.startswith("## "):
            head = line[3:].strip()
            if head.lower() == "abstract":
                section = "abstract"
                continue
            if head.lower().startswith("references"):
                section = "refs"
                continue
            section = "body"
            out["body"].append((S_H1, head))
            continue

        if line.startswith("### "):
            out["body"].append((S_H2, line[4:].strip()))
            continue
        if line.startswith("#### "):
            out["body"].append((S_H3, line[5:].strip()))
            continue

        if section == "abstract" and not line.startswith("**Keywords:**"):
            out["abstract"] = line
            continue

        if line.startswith("**Keywords:**"):
            out["keywords"] = line[len("**Keywords:**"):].strip()
            section = "body"
            continue

        if section == "refs":
            out["refs"].append(re.sub(r"^\d+\.\s*", "", line))   # style auto-numbers
            continue

        if re.match(r"^\*\*(%s)" % "|".join(BACK_MATTER_LABELS), line):
            out["back"].append(line)
            continue

        m = re.match(r"!\[\]\(([^)\s]+)", line)
        if m:
            out["body"].append(("image", m.group(1)))
            continue

        if re.match(r"\*\*Figure\s+S?\d+\.\*\*", line):
            out["body"].append((S_FIGCAP, line))
            continue

        if line.startswith("|"):
            out["body"].append((S_TEXT_NOIND, line))
            continue

        out["body"].append((S_TEXT, line))

    for k in ("title", "abstract", "keywords"):
        if not out[k]:
            raise SystemExit("markdown is missing the %s" % k)
    return out


# ---------------------------------------------------------------- docx surgery

def template_as_document(dst):
    """Copy the .dot to dst, re-declaring the main part as a document rather than a template."""
    src = zipfile.ZipFile(TEMPLATE)
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as out:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"wordprocessingml.template.main+xml",
                                    b"wordprocessingml.document.main+xml")
            out.writestr(item, data)
    src.close()


def find(paras, needle, start=0):
    """Index of the first paragraph whose text starts with needle. Loud on failure."""
    for i in range(start, len(paras)):
        if paras[i].text.strip().startswith(needle):
            return i
    raise SystemExit("template placeholder not found: %r — has the .dot changed?" % needle)


def retext(par, text):
    """Replace a template paragraph's text, keeping its style and paragraph properties."""
    for r in list(par.runs):
        r._r.getparent().remove(r._r)
    add_runs(par, text)


def insert_before(doc, anchor_el, style):
    """A new empty paragraph in the given style, placed immediately before anchor_el."""
    p = anchor_el.makeelement(qn("w:p"), {})
    anchor_el.addprevious(p)
    par = Paragraph(p, doc)
    par.style = doc.styles[style]
    return par


def drop(el):
    el.getparent().remove(el)


def drop_span(doc, first_par, stop_par):
    """Delete every body child from first_par up to (not including) stop_par.

    Ranges are cut at the body level rather than over doc.paragraphs, because the template's
    tutorial sections interleave example tables between their paragraphs; deleting only the
    paragraphs would leave those tables stranded in the finished manuscript.
    """
    body = doc.element.body
    children = list(body)
    a, b = children.index(first_par._p), children.index(stop_par._p)
    for el in children[a:b]:
        drop(el)


def build():
    md = parse_markdown()
    template_as_document(OUT)
    doc = Document(OUT)
    paras = doc.paragraphs

    # --- front matter: replace only the fields we own, leave affiliations/editor table alone
    retext(paras[find(paras, "Type of the Paper")], "Article")
    retext(paras[find(paras, "Title")], md["title"])
    retext(paras[find(paras, "A single paragraph of about 200 words")], md["abstract"])

    kw = paras[find(paras, "Keywords:")]
    for r in list(kw.runs)[1:]:      # keep the template's bold "Keywords: " label run
        r._r.getparent().remove(r._r)
    add_runs(kw, md["keywords"])

    # --- body: everything from the tutorial section through "6. Patents" is template filler
    i_start = find(paras, "0. How to Use This Template")
    i_end = find(paras, "Supplementary Materials:")
    anchor = paras[i_end]._p
    drop_span(doc, paras[i_start], paras[i_end])

    n_fig = 0
    for style, text in md["body"]:
        if style == "image":
            path = os.path.join(FIGDIR, text)
            par = insert_before(doc, anchor, S_FIG)
            if os.path.exists(path):
                w, h = Image.open(path).size
                width = min(FIG_W, FIG_H_MAX * w / h)
                par.add_run().add_picture(path, width=Inches(width))
                n_fig += 1
            else:
                par.add_run("[missing figure: %s]" % text)
        else:
            add_runs(insert_before(doc, anchor, style), text)

    # --- back matter: same set of statements, our text
    paras = doc.paragraphs
    i_back = find(paras, "Supplementary Materials:")
    i_abbr = find(paras, "Abbreviations", i_back)
    anchor = paras[i_abbr]._p
    drop_span(doc, paras[i_back], paras[i_abbr])
    for line in md["back"]:
        add_runs(insert_before(doc, anchor, S_BACK), line)

    # --- abbreviations: reuse the template's table, cloning its first row for extra entries
    table = None
    for t in doc.tables:
        if t.rows and t.rows[0].cells[0].text.strip() == "MDPI":
            table = t
            break
    if table is None:
        raise SystemExit("abbreviations table not found in the template")
    proto = table.rows[0]._tr
    while len(table.rows) > 1:
        drop(table.rows[-1]._tr)
    for _ in range(len(ABBREVIATIONS) - 1):
        import copy
        table._tbl.append(copy.deepcopy(proto))
    for row, (short, long) in zip(table.rows, ABBREVIATIONS):
        for cell, val in zip(row.cells, (short, long)):
            retext(cell.paragraphs[0], val)
            for extra in cell.paragraphs[1:]:
                drop(extra._p)

    # --- appendices: we have none, and empty template appendices must not ship
    paras = doc.paragraphs
    i_apx = find(paras, "Appendix A")
    i_refs = find(paras, "References", i_apx)
    drop_span(doc, paras[i_apx], paras[i_refs])

    # --- references: the style auto-numbers, so entries carry no "1." prefix
    paras = doc.paragraphs
    i_refs = find(paras, "References")
    i_note = find(paras, "Disclaimer/Publisher", i_refs)
    anchor = paras[i_note]._p
    drop_span(doc, paras[i_refs + 1], paras[i_note])
    for ref in md["refs"]:
        add_runs(insert_before(doc, anchor, S_REFS), ref)

    doc.save(OUT)

    from collections import Counter
    chk = Document(OUT)
    styles = Counter(p.style.name for p in chk.paragraphs)
    print("wrote %s  (%.1f MB)" % (OUT, os.path.getsize(OUT) / 1e6))
    print("  figures embedded %d · references %d · back matter %d"
          % (n_fig, styles.get(S_REFS, 0), styles.get(S_BACK, 0)))
    stray = {k: v for k, v in styles.items() if not k.startswith("MDPI") and k != "Normal"}
    print("  non-template styles:", stray or "none")
    with zipfile.ZipFile(OUT) as z:
        print("  zip integrity:", "OK" if z.testzip() is None else "BAD")


if __name__ == "__main__":
    build()
